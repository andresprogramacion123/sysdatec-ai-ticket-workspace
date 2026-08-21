import base64
import io
import logging
import mimetypes
from urllib.parse import urlparse

import openpyxl
import requests
from bs4 import BeautifulSoup
from docx import Document

logger = logging.getLogger(__name__)

REQUEST_TIMEOUT_SECONDS = 5
MAX_DOWNLOAD_BYTES = 10 * 1024 * 1024  # 10MB hard cap on what we'll download at all
MAX_MEDIA_BYTES = 5 * 1024 * 1024  # 5MB cap for PDF/image content sent to the API
MAX_TEXT_CHARS = 1500
DOWNLOAD_CHUNK_SIZE = 8192
REQUEST_HEADERS = {"User-Agent": "sysdatec-ai-ticket-workspace/1.0 (+attachment-parser)"}

DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
XLSX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

EXTENSION_CONTENT_TYPE_FALLBACKS = {
    ".html": "text/html",
    ".htm": "text/html",
    ".txt": "text/plain",
    ".pdf": "application/pdf",
    ".docx": DOCX_CONTENT_TYPE,
    ".xlsx": XLSX_CONTENT_TYPE,
}


def _download(url: str) -> tuple[bytes, str, str | None] | None:
    with requests.get(
        url, stream=True, timeout=REQUEST_TIMEOUT_SECONDS, headers=REQUEST_HEADERS
    ) as response:
        response.raise_for_status()

        raw_content_type = response.headers.get("Content-Type", "")
        content_type = raw_content_type.split(";")[0].strip().lower()
        charset = None
        for part in raw_content_type.split(";")[1:]:
            if "=" in part and part.strip().lower().startswith("charset="):
                charset = part.split("=", 1)[1].strip().strip('"')
                break

        chunks: list[bytes] = []
        total_bytes = 0
        for chunk in response.iter_content(chunk_size=DOWNLOAD_CHUNK_SIZE):
            total_bytes += len(chunk)
            if total_bytes > MAX_DOWNLOAD_BYTES:
                return None
            chunks.append(chunk)

        return b"".join(chunks), content_type, charset


def _resolve_content_type(content_type: str, url: str) -> str:
    if content_type and content_type != "application/octet-stream":
        return content_type

    extension = "." + urlparse(url).path.rsplit(".", 1)[-1].lower() if "." in urlparse(url).path else ""
    return EXTENSION_CONTENT_TYPE_FALLBACKS.get(extension, content_type)


def _extract_html_text(content: bytes, charset: str | None) -> str:
    soup = BeautifulSoup(content, "html.parser", from_encoding=charset)
    for tag in soup(["script", "style"]):
        tag.decompose()
    return soup.get_text(separator=" ", strip=True)


def _decode_text(content: bytes, charset: str | None) -> str:
    for encoding in filter(None, [charset, "utf-8"]):
        try:
            return content.decode(encoding)
        except (LookupError, UnicodeDecodeError):
            continue
    return content.decode("latin-1")


def _extract_docx_text(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text)


def _extract_xlsx_text(content: bytes) -> str:
    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        lines.append(f"[{sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            values = [str(cell) for cell in row if cell is not None]
            if values:
                lines.append(" | ".join(values))
    return "\n".join(lines)


def extract_attachment_content(attachment_url: str) -> dict[str, str] | None:
    try:
        downloaded = _download(attachment_url)
        if downloaded is None:
            return None
        content, content_type, charset = downloaded

        content_type = _resolve_content_type(content_type, attachment_url)

        if content_type in ("text/html", "text/plain"):
            if content_type == "text/html":
                text = _extract_html_text(content, charset)
            else:
                text = _decode_text(content, charset)
            return {"type": "text", "content": text[:MAX_TEXT_CHARS]}

        if content_type == DOCX_CONTENT_TYPE:
            text = _extract_docx_text(content)
            return {"type": "text", "content": text[:MAX_TEXT_CHARS]}

        if content_type == XLSX_CONTENT_TYPE:
            text = _extract_xlsx_text(content)
            return {"type": "text", "content": text[:MAX_TEXT_CHARS]}

        if content_type == "application/pdf":
            if len(content) > MAX_MEDIA_BYTES:
                return None
            return {
                "type": "pdf",
                "media_type": "application/pdf",
                "data": base64.b64encode(content).decode("ascii"),
            }

        if content_type.startswith("image/"):
            if len(content) > MAX_MEDIA_BYTES:
                return None
            media_type = content_type or (
                mimetypes.guess_type(attachment_url)[0] or "image/jpeg"
            )
            return {
                "type": "image",
                "media_type": media_type,
                "data": base64.b64encode(content).decode("ascii"),
            }

        return None
    except Exception:
        logger.warning(
            "Failed to extract attachment content from %s", attachment_url, exc_info=True
        )
        return None
